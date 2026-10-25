#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Created : 20 Oct 2026
Authors : bruce-vdb, claude-opus-5-5
Ported from : A_Template_AnalysisQuery_anthropic_v3.4.py

-------------------------------------------------
A_Template_AnalysisQuery_google-genai_v3.0.py
-------------------------------------------------
  * Send one or more files of ANY permitted type for analysis to Gemini
  * Stream the reply to a text file that includes a) analytical results, and
    b) the Python code to generate any requested graphics (plus the code
    Gemini actually ran in its sandbox, and that code's output)
  * Save any files (charts/images) Gemini generates
  * Clean up (remove from storage) the uploaded files


---------------------------------------------------------------------------
QUICK START
---------------------------------------------------------------------------
  1. Check packages installed in venv (ai-apis) and perform updates

<<< bash
    cd ~/spyder-6/envs && source ./ai-apis/bin/activate &&
    pip list &&

    python3 -m pip install --upgrade pip &&
    pip install --upgrade google-genai &&

    # optional -- only needed for "container" (converted) file types
    pip install --upgrade pandas openpyxl xlrd odfpy pyarrow numpy \
                          python-docx python-pptx pillow &&
    deactivate

>>>
     Optional system tool: LibreOffice (soffice) converts .docx .doc .odt
     .rtf .pptx .ppt .odp to PDF so Gemini sees layout, tables and images.
         sudo apt install libreoffice-core libreoffice-writer libreoffice-impress

  2. Set your API key as ENV VARIABLE (once):  GEMINI_API_KEY
     (GOOGLE_API_KEY is also accepted)
  3. Edit SECTION 1 below (files, prompt)

  4. Run file (runs in Spyder within virtual environment, ai-apis)
      * To run file w/o CLI toggles
          a) Click on the green Run file icon in Toolbar or press F5

      * To run file w/ CLI toggles
          a) place copy of file on Desktop
          b) run file using IPython interpreter
runfile('/home/bruce-vdb/Desktop/A_Template_AnalysisQuery_google-genai_v3.0.py', args='--dry-run')

  Command-line examples (CLI files replace INPUT_FILES for that run):
     ./A_Template_...py report.pdf notes.txt -p "Summarize these files"
     ./A_Template_...py photo.png diagram.webp -p "Describe these images"
     ./A_Template_...py talk.mp3 demo.mp4 -p "Summarize the audio and video"
     ./A_Template_...py -c sales.xlsx budget.docx -p "Compare the totals"
     ./A_Template_...py data.csv -c data.csv -p "..."   # same file, both routes
     ./A_Template_...py paper.pdf --code-exec on -p "Plot ... as a chart"
     ./A_Template_...py script.py --prompt-file my_prompt.txt --overwrite
     ./A_Template_...py --dry-run            # show routing, upload nothing
     ./A_Template_...py --list-types         # list supported extensions
     ./A_Template_...py --purge-all          # delete ALL Files-API files
     ./A_Template_...py --help


---------------------------------------------------------------------------
HOW FILES ARE SENT TO GEMINI  (chosen automatically from the extension)
---------------------------------------------------------------------------
  route       Gemini part            file types
  ---------   --------------------   -----------------------------------------
  document    file_data (Files API)  .pdf  and plain-text / source-code files
                                     (.txt .md .csv .py .json .html .sql ...)
  image       file_data              .jpg .jpeg .png .webp .heic .heif
                                     (.gif .bmp .tif .tiff .ico are converted
                                     to PNG locally first -- needs pillow)
  media       file_data              audio: .mp3 .wav .aiff .aac .ogg .flac
                                     video: .mp4 .mpeg .mov .avi .webm .wmv ...
  container   file_data + the        Gemini's sandbox cannot receive arbitrary
              code-execution tool    binaries, so these are CONVERTED LOCALLY:
                                       .xlsx .xls .ods       -> CSV per sheet
                                       .docx .pptx .odt ...  -> PDF (LibreOffice)
                                                                or plain text
                                       .parquet .feather     -> CSV
                                       .npy .npz             -> CSV
                                       .sqlite .db           -> CSV per table
                                       .epub                 -> text
                                       .zip .tar .gz .tgz    -> unpacked members
                                     The code-execution tool is enabled
                                     automatically (see CODE_EXECUTION).

---------------------------------------------------------------------------
DIFFERENCES FROM THE ANTHROPIC VERSION
---------------------------------------------------------------------------
  * No "pause_turn": Gemini finishes tool runs inside one call; only
    MAX_TOKENS triggers a continuation pass.
  * The Gemini sandbox is stateless between calls (no container reuse).
  * Generated files arrive as inline image bytes (matplotlib output), not as
    downloadable file ids.  Set GENERATED_IMAGE_FORMAT = "jpeg" to convert.
  * Uploaded files auto-expire after 48 h, but are still deleted at the end.
  * Uses the Gemini Developer API (Files API is not available on Vertex AI).

-----------
REVISIONS
-----------
v3.0 -- port of the Anthropic v3.4 template to google-genai / Gemini

------------
NOTES
-------------
**Getting the JPEG diagrams from your current prompt**

With only PDFs as input, `CODE_EXECUTION = "auto"` leaves the tool off, exactly like the v3.4 script. Gemini will then 
write the diagram code into the output file. To have Gemini draw the diagrams and save them as JPEG files:
1. Set `CODE_EXECUTION = "on"` (or run with `--code-exec on`).
2. Set `GENERATED_IMAGE_FORMAT = "jpeg"`. Gemini's sandbox returns matplotlib charts as PNG, and this setting converts 
them; it needs Pillow (`pillow`).

**Model and SDK**
- `gemini-3.8-flash` is used as you requested. If the API rejects it, change `MODEL` to an available ID.
- `thinking_level` requires a recent `google-genai` version. On older versions the script prints a note and continues 
without it.

**Small fixes carried over from v3.4**
- `DOC_DIR = Path.home() / "/home/bruce-vdb/Desktop"` only worked because joining an absolute path discards `Path.home()`. It is now `Path.home() / "Desktop"`.
- The prompt used backslash line continuations inside the string. These inserted runs of indentation spaces into the prompt, 
despite the "no stray spaces" comment. The prompt now uses adjacent string literals.


**Optional dependencies**
- These are needed only for container file types; `--list-types` shows which ones are missing.
- LibreOffice gives the best results for Word and PowerPoint files, because Gemini then receives a PDF with layout, 
tables and images intact.
- `.7z`, `.pkl` and `.h5` are not converted. Pickle files are skipped because loading them can execute code. 
Convert these files to CSV, text or PDF first.



"""

from __future__ import annotations

import argparse
import csv
import gzip
import importlib.util
import io
import mimetypes
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tarfile
import tempfile
import time
import zipfile
from dataclasses import dataclass
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable

try:
    from google import genai
    from google.genai import errors, types
except ImportError:
    sys.exit("The 'google-genai' package is not installed.  Run:  pip install --upgrade google-genai")


# =========================================================================
#  SECTION 1 -- USER SETTINGS  (edit these)
# =========================================================================

# ---- Directory that holds the input files and receives the output ---------
DOC_DIR = Path.home() / "Desktop"

# ---- Input files (on Desktop) ----------------------------------------------
#   * One entry per line; add, remove or comment out (#) lines freely.
#   * Names are relative to DOC_DIR, or give an absolute path.
#   * Entry formats:
#         "file.ext"                      -> route chosen from the extension
#         ("file.ext", "container")       -> force a route:
#                                            "document" | "image" | "media" | "container"
INPUT_FILES = [
    # ---- PDF documents ----------------------------------------------
     "Woods_Resilience.pdf",
     "Ashby_Resilience.pdf",

    # ---- Plain text / source code -----------------------------------
#    "A_Template_ResearchQuery_openai_v3.0.py",
#    "Errors.txt",

    # ---- Images -----------------------------------------------------
    # "Screenshot.png",

    # ---- Audio / video ----------------------------------------------
    # "Lecture.mp3",

    # ---- Other files (converted locally, analyzed with code) --------
    # "Sales_2025.xlsx",
]

# ---- EXAMPLES: one or more for every supported file type --------------
#   Copy any line into INPUT_FILES above (and change the name).
#
#   PDF  (route: document) ..............................................
#     "Annual_Report_2025.pdf",
#
#   PLAIN TEXT  (route: document; file must be UTF-8 text) ..............
#     "Notes.txt",              "README.md",            "Manual.rst",
#     "Data.csv",               "Data.tsv",             "Server.log",
#     "Config.json",            "Records.jsonl",        "Settings.yaml",
#     "Settings.yml",           "pyproject.toml",       "Setup.ini",
#     "Feed.xml",               "Page.html",            "Style.css",
#     "Paper.tex",              "Refs.bib",             "Subtitles.srt",
#     "Notebook.ipynb",         "Makefile",             "Dockerfile",
#
#   SOURCE CODE  (route: document) ......................................
#     "Script.py",   "App.js",   "App.ts",    "Main.java",  "Main.c",
#     "Main.cpp",    "Lib.h",    "Prog.cs",   "Main.go",    "Main.rs",
#     "Tool.rb",     "Index.php","Run.sh",    "Run.bat",    "Run.ps1",
#     "Query.sql",   "Model.R",  "Calc.m",    "App.swift",  "App.kt",
#
#   IMAGES  (route: image) ..............................................
#     "Photo.jpg",   "Photo.jpeg",   "Chart.png",   "Pic.webp",   "Phone.heic",
#     "Anim.gif",    "Bitmap.bmp",   "Scan.tiff"    (last three -> PNG locally)
#
#   AUDIO / VIDEO  (route: media) .......................................
#     "Talk.mp3",    "Memo.wav",     "Song.flac",   "Clip.mp4",   "Demo.mov",
#
#   OTHER FILES  (route: container -- converted locally, then code exec) .
#     "Sales_2025.xlsx",        "Legacy.xls",           "Sheet.ods",
#     "Contract.docx",          "Letter.odt",           "Deck.pptx",
#     "Archive.zip",            "Backup.tar.gz",        "Table.parquet",
#     "Inventory.sqlite",       "Arrays.npy",           "Ebook.epub",
#
#   FORCING A ROUTE  (tuple form) .......................................
#     ("Big_Dataset.csv", "container"),  # analyze with code, not read as text
#     ("Vector.svg", "document"),        # read SVG markup as text
#     ("Photo.png", "container"),        # process the image with Python

# ---- The prompt (adjacent string literals = no stray spaces) ----------
PROMPT = (
    "Review and analyze the file, Ashby_Resilience.pdf, which is an excerpt "
    "from Ashby's, Introduction to Cybernetics. Compare and contrast Ashby's "
    "concept of the Law of Requisite Variety with the four concepts in "
    "Woods's paper, Four concepts for resilience and the implications for "
    "the future of resilience engineering. Do the following: 1) Create a "
    "single diagram that frames The Law of Requisite Variety in terms of "
    "resilience, and 2) Create a single diagram that compares and contrasts "
    "The Law of Requisite Variety with Woods's concepts of resilience. The "
    "diagrams should be in .jpeg format."
)

# ---- Optional system prompt (None = no system prompt) -----------------
SYSTEM_PROMPT = (
    "You are an expert in sustainability, the analysis of complex adaptive "
    "systems, and cybernetics. When an additional file, chart, table, or "
    "script would be useful, create it with the code-execution tool (if it "
    "is available) so the user can download it; otherwise provide complete, "
    "runnable Python code that creates it."
)

# ---- Output -------------------------------------------------------------
OUTPUT_FILE = "QueryResponse_Gemini.txt"
APPEND_MODE = True               # True = append each run; False = overwrite
DOWNLOAD_GENERATED_FILES = True  # save charts/images Gemini creates in the sandbox
GENERATED_PREFIX = "generated_"  # prefix for those saved files
GENERATED_IMAGE_FORMAT = None    # None = keep as returned (PNG); or "jpeg", "png", "webp"

# ---- Code execution -----------------------------------------------------
#   "auto" = enable only when a container file is present (like the Anthropic script)
#   "on"   = always enable (Gemini can then RENDER requested charts itself)
#   "off"  = never enable
CODE_EXECUTION = "auto"

# ---- Unknown extensions -------------------------------------------------
#   "container" = send to the sandbox as text (file must be UTF-8 text)
#   "error"     = refuse unknown extensions
UNKNOWN_TYPE_POLICY = "container"

# ---- Model / output size -----------------------------------------------
MODEL = "gemini-3.8-flash"  # or another current model id
MAX_TOKENS = 65536          # max output tokens per pass (lower if rejected)
MAX_CONTINUATIONS = 3       # extra passes if the reply is cut off
THINKING_LEVEL = "high"     # "minimal" | "low" | "medium" | "high" | None (model default)


# =========================================================================
#  SECTION 2 -- FIXED SETTINGS  (normally no need to edit)
# =========================================================================

MAX_UPLOAD_BYTES = 2 * 1024 ** 3            # Files API per-file limit (2 GB)
MAX_PDF_BYTES = 50 * 1024 ** 2              # Gemini PDF limit (50 MB)
FILE_PROCESS_TIMEOUT = 600                  # seconds to wait for ACTIVE state
MAX_ARCHIVE_MEMBERS = 50                    # files taken from one archive

ROUTES = ("document", "image", "media", "container")

# Text-like files -> document part with MIME text/plain
TEXT_EXTENSIONS = {
    # prose / markup / data
    ".txt", ".text", ".md", ".markdown", ".rst", ".csv", ".tsv", ".log",
    ".json", ".jsonl", ".ndjson", ".xml", ".html", ".htm", ".css",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".env", ".properties",
    ".tex", ".bib", ".srt", ".vtt", ".ipynb", ".svg",
    # source code
    ".py", ".pyw", ".js", ".mjs", ".ts", ".tsx", ".jsx", ".java", ".c",
    ".h", ".cpp", ".cc", ".hpp", ".cs", ".go", ".rs", ".rb", ".php",
    ".pl", ".lua", ".sh", ".bash", ".zsh", ".bat", ".cmd", ".ps1", ".sql",
    ".r", ".m", ".swift", ".kt", ".scala", ".dart", ".vb", ".f90",
}
TEXT_FILENAMES = {"makefile", "dockerfile", "readme", "license", "changelog"}

# Images Gemini reads natively -> image part
IMAGE_TYPES = {
    ".jpg":  "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png":  "image/png",
    ".webp": "image/webp",
    ".heic": "image/heic",
    ".heif": "image/heif",
}
# Images converted locally to PNG first (route stays "image")
IMAGE_CONVERT = {
    ".gif":  "image/gif",
    ".bmp":  "image/bmp",
    ".tif":  "image/tiff",
    ".tiff": "image/tiff",
    ".ico":  "image/x-icon",
}

# Audio / video -> media part
AUDIO_TYPES = {
    ".mp3": "audio/mp3", ".wav": "audio/wav", ".aiff": "audio/aiff",
    ".aif": "audio/aiff", ".aac": "audio/aac", ".ogg": "audio/ogg",
    ".flac": "audio/flac",
}
VIDEO_TYPES = {
    ".mp4": "video/mp4", ".mpeg": "video/mpeg", ".mpg": "video/mpg",
    ".mov": "video/mov", ".avi": "video/avi", ".flv": "video/x-flv",
    ".webm": "video/webm", ".wmv": "video/wmv", ".3gp": "video/3gpp",
}
MEDIA_TYPES = {**AUDIO_TYPES, **VIDEO_TYPES}

# Binary formats -> converted locally, sent as container files
#   ext: (original mime, converter key -- see CONVERTERS in SECTION 3)
CONTAINER_TYPES = {
    ".xlsx":    ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "spreadsheet"),
    ".xlsm":    ("application/vnd.ms-excel.sheet.macroEnabled.12", "spreadsheet"),
    ".xls":     ("application/vnd.ms-excel", "spreadsheet"),
    ".ods":     ("application/vnd.oasis.opendocument.spreadsheet", "spreadsheet"),
    ".docx":    ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "office"),
    ".doc":     ("application/msword", "office"),
    ".odt":     ("application/vnd.oasis.opendocument.text", "office"),
    ".rtf":     ("application/rtf", "office"),
    ".pptx":    ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "office"),
    ".ppt":     ("application/vnd.ms-powerpoint", "office"),
    ".odp":     ("application/vnd.oasis.opendocument.presentation", "office"),
    ".epub":    ("application/epub+zip", "epub"),
    ".zip":     ("application/zip", "archive"),
    ".tar":     ("application/x-tar", "archive"),
    ".gz":      ("application/gzip", "archive"),
    ".tgz":     ("application/gzip", "archive"),
    ".parquet": ("application/vnd.apache.parquet", "table"),
    ".feather": ("application/octet-stream", "table"),
    ".npy":     ("application/octet-stream", "numpy"),
    ".npz":     ("application/octet-stream", "numpy"),
    ".sqlite":  ("application/vnd.sqlite3", "sqlite"),
    ".sqlite3": ("application/vnd.sqlite3", "sqlite"),
    ".db":      ("application/vnd.sqlite3", "sqlite"),
}

# pip package names for modules used by the converters
PIP_NAMES = {"PIL": "pillow", "docx": "python-docx", "pptx": "python-pptx",
             "odf": "odfpy"}

# file extensions for generated (inline) files
GENERATED_EXT = {
    "image/png": ".png", "image/jpeg": ".jpeg", "image/webp": ".webp",
    "image/gif": ".gif", "image/svg+xml": ".svg", "text/csv": ".csv",
    "text/plain": ".txt", "application/pdf": ".pdf",
}

CONTINUE_PROMPT = (
    "Your previous reply was cut off by the output-token limit. "
    "Continue exactly where you stopped -- no preamble, no repeated text."
)


# =========================================================================
#  SECTION 3 -- FUNCTIONS
# =========================================================================

@dataclass
class InputFile:
    path: Path
    route: str                  # "document" | "image" | "media" | "container"
    mime: str
    convert: str | None = None  # converter key, or None = send as-is


@dataclass
class UploadedFile:
    route: str
    label: str                  # name shown to Gemini
    mime: str
    name: str                   # Files API resource name ("files/abc123")
    uri: str
    converted_from: str | None = None


def enum_name(value) -> str:
    """Readable name of an SDK enum value (or 'NONE')."""
    if value is None:
        return "NONE"
    return getattr(value, "name", str(value))


def safe_name(name) -> str:
    return re.sub(r"[^\w.-]+", "_", str(name)).strip("._") or "item"


def text_mime(path: Path) -> str:
    return "text/csv" if path.suffix.lower() == ".csv" else "text/plain"


# ---- local converters ------------------------------------------------------
#   Each converter: func(path, workdir) -> [(local_path, mime, label), ...]
#                   check(path)         -> problem string, or None if usable

def missing_modules(*modules: str) -> str | None:
    missing = [m for m in modules if importlib.util.find_spec(m) is None]
    if not missing:
        return None
    return "needs:  pip install " + " ".join(PIP_NAMES.get(m, m) for m in missing)


def find_soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


# spreadsheets -> one CSV per sheet
SPREADSHEET_ENGINES = {".xlsx": "openpyxl", ".xlsm": "openpyxl",
                       ".xls": "xlrd", ".ods": "odf"}


def check_spreadsheet(path: Path) -> str | None:
    return missing_modules("pandas",
                           SPREADSHEET_ENGINES.get(path.suffix.lower(), "openpyxl"))


def conv_spreadsheet(path: Path, work: Path) -> list:
    import pandas as pd
    outs = []
    for sheet, df in pd.read_excel(path, sheet_name=None).items():
        out = work / f"{path.stem}__{safe_name(sheet)}.csv"
        df.to_csv(out, index=False)
        outs.append((out, "text/csv", f"{path.name} [sheet: {sheet}]"))
    if not outs:
        raise RuntimeError("workbook has no sheets")
    return outs


# parquet / feather -> CSV
def check_table(path: Path) -> str | None:
    return missing_modules("pandas", "pyarrow")


def conv_table(path: Path, work: Path) -> list:
    import pandas as pd
    df = (pd.read_parquet(path) if path.suffix.lower() == ".parquet"
          else pd.read_feather(path))
    out = work / f"{path.stem}.csv"
    df.to_csv(out, index=False)
    return [(out, "text/csv", f"{path.name} (as CSV)")]


# numpy -> CSV (arrays with >2 dims are flattened to 2-D; shape kept in header)
def check_numpy(path: Path) -> str | None:
    return missing_modules("numpy")


def conv_numpy(path: Path, work: Path) -> list:
    import numpy as np
    loaded = np.load(path, allow_pickle=False)
    try:
        if hasattr(loaded, "files"):                  # .npz archive
            arrays = {k: loaded[k] for k in loaded.files}
        else:
            arrays = {path.stem: loaded}
    finally:
        if hasattr(loaded, "close"):
            loaded.close()
    outs = []
    for key, arr in arrays.items():
        arr = np.atleast_1d(arr)
        shape, dtype = arr.shape, arr.dtype
        if arr.ndim > 2:
            arr = arr.reshape(arr.shape[0], -1)
        out = work / f"{path.stem}__{safe_name(key)}.csv"
        np.savetxt(out, arr, delimiter=",", fmt="%s",
                   header=f"original shape {shape}, dtype {dtype}")
        outs.append((out, "text/csv", f"{path.name} [array: {key}]"))
    return outs


# sqlite -> one CSV per table/view
def conv_sqlite(path: Path, work: Path) -> list:
    outs = []
    con = sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True)
    try:
        tables = [r[0] for r in con.execute(
            "SELECT name FROM sqlite_master WHERE type IN ('table', 'view') "
            "AND name NOT LIKE 'sqlite_%' ORDER BY name")]
        for table in tables:
            quoted = '"' + table.replace('"', '""') + '"'
            cur = con.execute(f"SELECT * FROM {quoted}")
            out = work / f"{path.stem}__{safe_name(table)}.csv"
            with open(out, "w", newline="", encoding="utf-8") as fh:
                writer = csv.writer(fh)
                writer.writerow([d[0] for d in cur.description])
                writer.writerows(cur)
            outs.append((out, "text/csv", f"{path.name} [table: {table}]"))
    finally:
        con.close()
    if not outs:
        raise RuntimeError("database has no tables")
    return outs


# Word / PowerPoint / ODF / RTF -> PDF (LibreOffice) or plain text
OFFICE_TEXT_FALLBACK = {".docx": "docx", ".pptx": "pptx"}


def check_office(path: Path) -> str | None:
    if find_soffice():
        return None
    mod = OFFICE_TEXT_FALLBACK.get(path.suffix.lower())
    if mod and not missing_modules(mod):
        return None
    hint = f"  (or: pip install {PIP_NAMES[mod]})" if mod else ""
    return "needs LibreOffice (soffice) to convert it to PDF" + hint


def docx_text(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for i, table in enumerate(d.tables, 1):
        lines.append(f"\n[Table {i}]")
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def pptx_text(path: Path) -> str:
    from pptx import Presentation
    lines = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        lines.append(f"\n--- Slide {i} ---")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                lines.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def conv_office(path: Path, work: Path) -> list:
    soffice = find_soffice()
    if soffice:
        profile = (work / "lo_profile").resolve().as_uri()  # avoids clash with an open LibreOffice
        cmd = [soffice, f"-env:UserInstallation={profile}", "--headless",
               "--convert-to", "pdf", "--outdir", str(work), str(path)]
        subprocess.run(cmd, check=True, capture_output=True, timeout=300)
        pdf = work / f"{path.stem}.pdf"
        if not pdf.is_file():
            raise RuntimeError("LibreOffice produced no PDF")
        return [(pdf, "application/pdf", f"{path.name} (as PDF)")]
    text = docx_text(path) if path.suffix.lower() == ".docx" else pptx_text(path)
    out = work / f"{path.stem}.txt"
    out.write_text(text, encoding="utf-8")
    return [(out, "text/plain", f"{path.name} (as text)")]


# epub -> plain text
class _TextExtractor(HTMLParser):
    BREAKS = {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self):
        super().__init__()
        self.chunks: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in self.BREAKS:
            self.chunks.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.chunks.append(data)

    def text(self) -> str:
        return re.sub(r"\n\s*\n+", "\n\n", "".join(self.chunks)).strip()


def conv_epub(path: Path, work: Path) -> list:
    sections = []
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist()
                       if n.lower().endswith((".xhtml", ".html", ".htm")))
        for name in names:
            parser = _TextExtractor()
            parser.feed(z.read(name).decode("utf-8", errors="replace"))
            text = parser.text()
            if text:
                sections.append(f"\n\n=== {name} ===\n{text}")
    if not sections:
        raise RuntimeError("no readable text in EPUB")
    out = work / f"{path.stem}.txt"
    out.write_text("".join(sections).strip(), encoding="utf-8")
    return [(out, "text/plain", f"{path.name} (as text)")]


# gif / bmp / tiff / ico -> PNG
def check_image(path: Path) -> str | None:
    return missing_modules("PIL")


def conv_image(path: Path, work: Path) -> list:
    from PIL import Image
    with Image.open(path) as img:
        img.seek(0)                                   # first frame of GIF/TIFF
        frame = (img.copy() if img.mode in ("RGB", "RGBA", "L", "LA")
                 else img.convert("RGBA"))
    out = work / f"{path.stem}.png"
    frame.save(out, "PNG")
    return [(out, "image/png", f"{path.name} (as PNG)")]


# archives -> readable members (converted where needed; no nested archives)
def conv_archive(path: Path, work: Path) -> list:
    dest = work / "unpacked"
    dest.mkdir(parents=True, exist_ok=True)
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            z.extractall(dest)                        # zipfile sanitizes paths
    elif tarfile.is_tarfile(path):
        with tarfile.open(path) as t:
            if hasattr(tarfile, "data_filter"):
                t.extractall(dest, filter="data")     # safe extraction
            else:
                t.extractall(dest)
    elif path.name.lower().endswith(".gz"):
        with gzip.open(path, "rb") as src, open(dest / path.stem, "wb") as dst:
            shutil.copyfileobj(src, dst)
    else:
        raise RuntimeError("unrecognized archive format")

    outs, skipped = [], []
    for member in sorted(p for p in dest.rglob("*") if p.is_file()):
        if len(outs) >= MAX_ARCHIVE_MEMBERS:
            print(f"  archive {path.name}: stopped at {MAX_ARCHIVE_MEMBERS} files")
            break
        rel = member.relative_to(dest).as_posix()
        label = f"{path.name}/{rel}"
        route, mime, conv = detect_route(member)
        try:
            if route in ("document", "image") and conv is None:
                if mime == "text/plain":
                    if check_text_file(member):
                        skipped.append(rel)
                        continue
                    mime = text_mime(member)
                outs.append((member, mime, label))
            elif conv and conv != "archive" and CONVERTERS[conv].check(member) is None:
                sub = work / f"member_{len(outs):03d}"
                sub.mkdir()
                for p, m, lab in CONVERTERS[conv].func(member, sub):
                    tail = lab[len(member.name):] if lab.startswith(member.name) else f" -> {lab}"
                    outs.append((p, m, label + tail))
            elif route is None and check_text_file(member) is None:
                outs.append((member, "text/plain", label))
            else:
                skipped.append(rel)
        except Exception:
            skipped.append(rel)
    if skipped:
        shown = ", ".join(skipped[:10]) + (" ..." if len(skipped) > 10 else "")
        print(f"  archive {path.name}: skipped {len(skipped)} member(s): {shown}")
    if not outs:
        raise RuntimeError("archive contains no files Gemini can read")
    return outs


@dataclass(frozen=True)
class Converter:
    func: Callable[[Path, Path], list]
    check: Callable[[Path], "str | None"]
    describe: str


CONVERTERS: dict[str, Converter] = {
    "spreadsheet": Converter(conv_spreadsheet, check_spreadsheet, "each sheet -> CSV"),
    "table":       Converter(conv_table, check_table, "-> CSV"),
    "numpy":       Converter(conv_numpy, check_numpy, "each array -> CSV"),
    "sqlite":      Converter(conv_sqlite, lambda p: None, "each table -> CSV"),
    "office":      Converter(conv_office, check_office, "-> PDF (LibreOffice) or text"),
    "epub":        Converter(conv_epub, lambda p: None, "-> text"),
    "image":       Converter(conv_image, check_image, "-> PNG"),
    "archive":     Converter(conv_archive, lambda p: None, "unpack -> readable members"),
}


# ---- command line --------------------------------------------------------

def parse_args() -> argparse.Namespace:
    """Command-line options; every option defaults to the SECTION 1 value."""
    p = argparse.ArgumentParser(
        description="Send files of any permitted type + a prompt to Gemini "
                    "and save the streamed reply.",
    )
    p.add_argument("files", nargs="*",
                   help="input files, route chosen by extension "
                        "(default: INPUT_FILES in the script)")
    p.add_argument("-c", "--container", nargs="+", default=[], metavar="FILE",
                   help="files to force into the code-execution route")
    p.add_argument("-d", "--dir", type=Path, default=DOC_DIR,
                   help=f"folder for input/output files (default: {DOC_DIR})")
    p.add_argument("-p", "--prompt", default=None,
                   help="prompt text (default: PROMPT in the script)")
    p.add_argument("--prompt-file", type=Path, default=None,
                   help="read the prompt from a text file")
    p.add_argument("-s", "--system", default=SYSTEM_PROMPT,
                   help="optional system prompt")
    p.add_argument("-o", "--output", default=OUTPUT_FILE,
                   help=f"output file name (default: {OUTPUT_FILE})")
    p.add_argument("--overwrite", action="store_true",
                   help="overwrite the output file instead of appending")
    p.add_argument("-m", "--model", default=MODEL,
                   help=f"model id (default: {MODEL})")
    p.add_argument("--max-tokens", type=int, default=MAX_TOKENS,
                   help=f"max output tokens per pass (default: {MAX_TOKENS})")
    p.add_argument("--max-continuations", type=int, default=MAX_CONTINUATIONS,
                   help=f"extra passes if cut off (default: {MAX_CONTINUATIONS})")
    p.add_argument("--code-exec", choices=("auto", "on", "off"), default=CODE_EXECUTION,
                   help=f"enable the code-execution tool (default: {CODE_EXECUTION})")
    p.add_argument("--thinking", choices=("minimal", "low", "medium", "high", "none"),
                   default=THINKING_LEVEL or "none",
                   help=f"thinking level (default: {THINKING_LEVEL or 'none'})")
    p.add_argument("--dry-run", action="store_true",
                   help="show how each file would be sent, then exit (no upload)")
    p.add_argument("--list-types", action="store_true",
                   help="list supported file extensions, then exit")
    p.add_argument("--purge-all", action="store_true",
                   help="delete EVERY file in your Gemini Files storage, then exit")
    args = p.parse_args()
    args.dir = args.dir.expanduser()
    return args


def list_types() -> None:
    """Print every supported extension grouped by route."""
    def show(title: str, exts) -> None:
        print(f"\n{title}")
        print("  " + "  ".join(sorted(exts)))
    show("document (PDF):", [".pdf"])
    show("document (plain text, must be UTF-8):", TEXT_EXTENSIONS)
    print("  also extension-less: " + ", ".join(sorted(TEXT_FILENAMES)))
    show("image (native):", IMAGE_TYPES)
    show("image (converted locally to PNG):", IMAGE_CONVERT)
    issue = check_image(Path("x.gif"))
    if issue:
        print(f"  [{issue}]")
    show("media (audio):", AUDIO_TYPES)
    show("media (video):", VIDEO_TYPES)

    print("\ncontainer (converted locally, then analyzed with the code-execution tool):")
    groups: dict[str, list[str]] = {}
    for ext, (_, conv) in CONTAINER_TYPES.items():
        groups.setdefault(conv, []).append(ext)
    for conv, exts in groups.items():
        c = CONVERTERS[conv]
        print(f"  {' '.join(sorted(exts)):42} {c.describe}")
        issues = {c.check(Path(f"x{e}")) for e in exts} - {None}
        for issue in sorted(issues):
            print(f"  {'':42} [{issue}]")
    if UNKNOWN_TYPE_POLICY == "container":
        print("\n  ...plus ANY other extension that contains UTF-8 text "
              "(UNKNOWN_TYPE_POLICY = 'container')")
    print()


# ---- client / prompt -----------------------------------------------------

def make_client() -> genai.Client:
    """Create the API client (key read from GEMINI_API_KEY or GOOGLE_API_KEY)."""
    api_key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
    if not api_key:
        sys.exit("ERROR: environment variable GEMINI_API_KEY (or GOOGLE_API_KEY) is not set.")
    return genai.Client(api_key=api_key)


def resolve_prompt(args: argparse.Namespace) -> str:
    """Prompt priority: --prompt-file, then --prompt, then PROMPT."""
    if args.prompt_file:
        path = args.prompt_file.expanduser()
        if not path.is_absolute():
            path = args.dir / path
        if not path.is_file():
            sys.exit(f"ERROR: prompt file not found: {path}")
        return path.read_text(encoding="utf-8").strip()
    return args.prompt if args.prompt else PROMPT


# ---- input files ---------------------------------------------------------

def detect_route(path: Path) -> tuple[str | None, str | None, str | None]:
    """(route, mime, converter) from the file extension; all None if unknown."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return "document", "application/pdf", None
    if ext in TEXT_EXTENSIONS or (not ext and path.name.lower() in TEXT_FILENAMES):
        return "document", "text/plain", None
    if ext in IMAGE_TYPES:
        return "image", IMAGE_TYPES[ext], None
    if ext in IMAGE_CONVERT:
        return "image", IMAGE_CONVERT[ext], "image"
    if ext in MEDIA_TYPES:
        return "media", MEDIA_TYPES[ext], None
    if ext in CONTAINER_TYPES:
        mime, conv = CONTAINER_TYPES[ext]
        return "container", mime, conv
    return None, None, None


def check_text_file(path: Path) -> str | None:
    """Return a problem description if the file is not usable as text."""
    data = path.read_bytes()
    if not data.strip():
        return "file is empty"
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        return "not UTF-8 text -- re-save it as UTF-8"
    return None


def normalize_entries(entries) -> list[tuple[str, str]]:
    """Accept "name" or ("name", "route") and return (name, route) pairs."""
    pairs = []
    for entry in entries:
        if isinstance(entry, (tuple, list)):
            if len(entry) != 2:
                sys.exit(f"ERROR: bad INPUT_FILES entry {entry!r}; "
                         'use "name" or ("name", "route")')
            pairs.append((str(entry[0]), str(entry[1]).strip().lower()))
        else:
            pairs.append((str(entry), "auto"))
    return pairs


def resolve_files(entries, doc_dir: Path) -> list[InputFile]:
    """Turn entries into InputFile objects; stop on any problem."""
    files, problems = [], []
    for name, route in normalize_entries(entries):
        path = Path(name).expanduser()
        if not path.is_absolute():
            path = doc_dir / path

        if not path.is_file():
            problems.append(f"  not found:            {path}")
            continue
        size = path.stat().st_size
        if size > MAX_UPLOAD_BYTES:
            problems.append(f"  larger than 2 GB:     {path}")
            continue
        if route not in ("auto",) + ROUTES:
            problems.append(f"  unknown route '{route}' (use {', '.join(ROUTES)}): {path}")
            continue

        auto_route, auto_mime, auto_conv = detect_route(path)
        text_checked = False

        if route == "auto":
            if auto_route is None:
                if UNKNOWN_TYPE_POLICY != "container":
                    problems.append(f"  unsupported type:     {path}")
                    continue
                route = "container"
            else:
                route = auto_route

        if route == "document":
            if auto_mime not in ("application/pdf", "text/plain"):
                problems.append(f"  not PDF/text, cannot be a document: {path}")
                continue
            mime, conv = auto_mime, None
        elif route == "image":
            if auto_route != "image":
                problems.append(f"  not an image type, cannot be an image: {path}")
                continue
            mime, conv = auto_mime, auto_conv
        elif route == "media":
            if auto_route != "media":
                problems.append(f"  not audio/video, cannot be media: {path}")
                continue
            mime, conv = auto_mime, None
        else:  # container
            if auto_route == "media":
                problems.append(f"  audio/video cannot be a container file: {path}")
                continue
            if auto_route is None:
                issue = check_text_file(path)
                if issue:
                    problems.append(f"  unknown type and {issue}; convert it to "
                                    f"PDF/CSV/text first: {path}")
                    continue
                mime, conv, text_checked = text_mime(path), None, True
            else:
                mime = text_mime(path) if auto_mime == "text/plain" else auto_mime
                conv = auto_conv

        if conv:
            issue = CONVERTERS[conv].check(path)
            if issue:
                problems.append(f"  {issue}: {path}")
                continue
        elif mime.startswith("text/") and not text_checked:
            issue = check_text_file(path)
            if issue:
                problems.append(f"  {issue}: {path}")
                continue
        if mime == "application/pdf" and size > MAX_PDF_BYTES:
            problems.append(f"  PDF larger than 50 MB: {path}")
            continue

        files.append(InputFile(path, route, mime, conv))

    if problems:
        sys.exit("ERROR: problem(s) with input files:\n" + "\n".join(problems))
    if not files:
        sys.exit("ERROR: no input files given.")
    return files


def code_exec_enabled(setting: str, files: list[InputFile]) -> bool:
    if setting == "on":
        return True
    if setting == "off":
        return False
    return any(f.route == "container" for f in files)


def show_routing(files: list[InputFile], use_code_execution: bool) -> None:
    """Print a table of how each file will be sent."""
    print(f"\n{'file':36} {'route':10} {'mime':30} conversion")
    print(f"{'-' * 36} {'-' * 10} {'-' * 30} {'-' * 26}")
    for f in files:
        conv = CONVERTERS[f.convert].describe if f.convert else "-"
        print(f"{f.path.name[:36]:36} {f.route:10} {f.mime[:30]:30} {conv}")
    print(f"\ncode-execution tool: {'ENABLED' if use_code_execution else 'disabled'}")
    print()


# ---- upload / request ----------------------------------------------------

def wait_until_active(client: genai.Client, f):
    """Poll an uploaded file until Gemini has finished processing it."""
    start = time.monotonic()
    while enum_name(f.state) == "PROCESSING":
        if time.monotonic() - start > FILE_PROCESS_TIMEOUT:
            raise RuntimeError(f"timed out waiting for {f.name} to be processed")
        time.sleep(2)
        f = client.files.get(name=f.name)
    if enum_name(f.state) == "FAILED":
        raise RuntimeError(f"Gemini could not process {f.display_name} ({f.name})")
    return f


def upload_files(client: genai.Client, files: list[InputFile], workdir: Path,
                 uploaded: list[UploadedFile]) -> None:
    """Convert (if needed) and upload each file. Successes are appended to
    `uploaded` immediately, so the caller can clean up if a later step fails."""
    for idx, item in enumerate(files, 1):
        if item.convert:
            sub = workdir / f"{idx:02d}_{safe_name(item.path.stem)}"
            sub.mkdir(parents=True, exist_ok=True)
            conv = CONVERTERS[item.convert]
            try:
                outputs = conv.func(item.path, sub)
            except Exception as e:
                raise RuntimeError(f"could not convert {item.path.name}: {e}") from e
            print(f"Converted {item.path.name} -> {len(outputs)} file(s) [{conv.describe}]")
            source = item.path.name
        else:
            outputs = [(item.path, item.mime, item.path.name)]
            source = None

        for local, mime, label in outputs:
            if local.stat().st_size > MAX_UPLOAD_BYTES:
                raise RuntimeError(f"{label} is larger than 2 GB after conversion")
            f = client.files.upload(
                file=str(local),
                config=types.UploadFileConfig(mime_type=mime, display_name=label[:500]),
            )
            up = UploadedFile(item.route, label, mime, f.name, f.uri, source)
            uploaded.append(up)
            f = wait_until_active(client, f)
            up.uri = f.uri or up.uri
            print(f"Uploaded {label} [{item.route}] -> {f.name}")


def build_contents(uploaded: list[UploadedFile], prompt: str) -> list:
    """Labeled files first; prompt last (Gemini's recommended order)."""
    kinds = {"document": "Document", "image": "Image",
             "media": "Media", "container": "Code-execution file"}
    parts = []
    container_names = []
    for u in uploaded:
        label = f"{kinds[u.route]}: {u.label}"
        if u.converted_from:
            label += f"  (converted locally from {u.converted_from})"
        parts.append(types.Part.from_text(text=label))
        parts.append(types.Part.from_uri(file_uri=u.uri, mime_type=u.mime))
        if u.route == "container":
            container_names.append(u.label)

    if container_names:
        parts.append(types.Part.from_text(text=(
            "The following files are provided for analysis with your "
            "code-execution tool; load and process them with Python as "
            "needed: " + ", ".join(container_names))))
    parts.append(types.Part.from_text(text=prompt))
    return [types.Content(role="user", parts=parts)]


def build_config(*, max_tokens: int, system: str | None,
                 use_code_execution: bool, thinking_level: str | None):
    kwargs = {"max_output_tokens": max_tokens}
    if system:
        kwargs["system_instruction"] = system
    if use_code_execution:
        kwargs["tools"] = [types.Tool(code_execution=types.ToolCodeExecution())]
    if thinking_level:
        try:
            kwargs["thinking_config"] = types.ThinkingConfig(thinking_level=thinking_level)
        except Exception as e:  # older SDK without thinking_level
            print(f"Note: thinking_level ignored ({e}); "
                  "run  pip install --upgrade google-genai")
    return types.GenerateContentConfig(**kwargs)


def _is_plain_text(part) -> bool:
    return part.text is not None and not part.thought


def handle_part(part, out, history: list, generated: list) -> None:
    """Write one streamed part to `out`; keep a compact copy for continuations."""
    if part.thought:                                  # thought summaries: not saved
        return
    if part.executable_code:
        code = (part.executable_code.code or "").rstrip()
        out.write("\n\n[... Gemini is running code in the sandbox ...]\n"
                  f"```python\n{code}\n```\n")
        history.append(part)
    elif part.code_execution_result:
        r = part.code_execution_result
        out.write(f"[sandbox result: {enum_name(r.outcome)}]\n"
                  f"```\n{(r.output or '').rstrip()}\n```\n\n")
        history.append(part)
    elif part.inline_data and part.inline_data.data:
        mime = part.inline_data.mime_type or "application/octet-stream"
        generated.append((part.inline_data.data, mime))
        out.write(f"\n[generated file #{len(generated)}: {mime}]\n")
        # image bytes are not resent in continuation passes
    elif part.text is not None:
        out.write(part.text)
        prev = history[-1] if history else None
        if (prev is not None and _is_plain_text(prev)
                and not prev.thought_signature and not part.thought_signature):
            history[-1] = types.Part(text=prev.text + part.text)
        elif part.text or part.thought_signature:
            history.append(types.Part(text=part.text,
                                      thought_signature=part.thought_signature))
    out.flush()                                       # nothing lost if the run dies


def run_query(client: genai.Client, contents: list, out, generated: list, *,
              model: str, max_tokens: int, max_continuations: int,
              system: str | None, use_code_execution: bool,
              thinking_level: str | None) -> None:
    """Stream the reply to `out`. Continues automatically if the reply is cut
    off (MAX_TOKENS). Generated images are appended to `generated`."""
    config = build_config(max_tokens=max_tokens, system=system,
                          use_code_execution=use_code_execution,
                          thinking_level=thinking_level)

    for attempt in range(max_continuations + 1):
        history: list = []
        finish = None
        usage = None
        block_reason = None

        stream = client.models.generate_content_stream(
            model=model, contents=contents, config=config)
        for chunk in stream:
            if chunk.usage_metadata:
                usage = chunk.usage_metadata
            feedback = getattr(chunk, "prompt_feedback", None)
            if feedback and feedback.block_reason:
                block_reason = enum_name(feedback.block_reason)
            if not chunk.candidates:
                continue
            cand = chunk.candidates[0]
            if cand.finish_reason:
                finish = cand.finish_reason
            parts = cand.content.parts if cand.content and cand.content.parts else []
            for part in parts:
                handle_part(part, out, history, generated)

        fname = enum_name(finish)
        print(f"Pass {attempt + 1}: finish_reason = {fname}, "
              f"input tokens = {getattr(usage, 'prompt_token_count', None)}, "
              f"output tokens = {getattr(usage, 'candidates_token_count', None)}, "
              f"thinking tokens = {getattr(usage, 'thoughts_token_count', None)}")

        if block_reason:
            warning = f"\n\n[WARNING: prompt was blocked (block_reason = {block_reason}).]\n"
            out.write(warning)
            print(warning)
            return

        if fname != "MAX_TOKENS":
            if fname not in ("STOP", "NONE", "FINISH_REASON_UNSPECIFIED"):
                warning = f"\n\n[WARNING: reply ended early (finish_reason = {fname}).]\n"
                out.write(warning)
                print(warning)
            return                                    # finished

        if attempt == max_continuations:              # out of retries
            warning = ("\n\n[WARNING: reply still incomplete "
                       f"(finish_reason = {fname}) after all continuation "
                       "passes. Raise MAX_TOKENS or MAX_CONTINUATIONS.]\n")
            out.write(warning)
            print(warning)
            return

        if not history:
            print("Reply was empty (output budget used by thinking?) -- cannot continue.")
            return
        contents.append(types.Content(role="model", parts=history))
        contents.append(types.Content(role="user",
                                      parts=[types.Part.from_text(text=CONTINUE_PROMPT)]))
        print("Reply hit MAX_TOKENS -- requesting continuation...")


# ---- generated files / cleanup ------------------------------------------

def convert_image_bytes(data: bytes, fmt: str) -> tuple[bytes, str]:
    """Convert image bytes to another format (needs pillow)."""
    from PIL import Image
    fmt = "jpeg" if fmt.lower() in ("jpg", "jpeg") else fmt.lower()
    with Image.open(io.BytesIO(data)) as src:
        img = src
        if fmt == "jpeg" and img.mode in ("RGBA", "LA", "P"):
            rgba = img.convert("RGBA")
            img = Image.new("RGB", rgba.size, "white")
            img.paste(rgba, mask=rgba.getchannel("A"))
        elif fmt == "jpeg" and img.mode != "RGB":
            img = img.convert("RGB")
        buf = io.BytesIO()
        img.save(buf, fmt.upper(), **({"quality": 95} if fmt == "jpeg" else {}))
    return buf.getvalue(), f"image/{fmt}"


def save_generated_files(generated: list, out_dir: Path,
                         convert_to: str | None = None) -> None:
    """Write the files (images) Gemini created to out_dir."""
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    for n, (data, mime) in enumerate(generated, 1):
        try:
            if convert_to and mime.startswith("image/"):
                try:
                    data, mime = convert_image_bytes(data, convert_to)
                except Exception as e:
                    print(f"Could not convert generated file #{n} to {convert_to} "
                          f"({e}); saving original.")
            ext = GENERATED_EXT.get(mime) or mimetypes.guess_extension(mime) or ".bin"
            dest = out_dir / f"{GENERATED_PREFIX}{stamp}_{n:02d}{ext}"
            dest.write_bytes(data)
            print(f"Saved generated file -> {dest}")
        except Exception as e:
            print(f"Could not save generated file #{n}: {e}")


def delete_uploads(client: genai.Client, uploaded: list[UploadedFile]) -> None:
    """Delete just the files uploaded during this run."""
    for u in uploaded:
        try:
            client.files.delete(name=u.name)
            print(f"Deleted {u.label} ({u.name})")
        except Exception as e:
            print(f"Could not delete {u.label} ({u.name}): {e}")


def purge_workspace(client: genai.Client) -> None:
    """Delete EVERY file in the Gemini Files storage (asks for confirmation)."""
    files = list(client.files.list())
    if not files:
        print("Files storage is already empty.")
        return
    answer = input(f"Delete ALL {len(files)} file(s) in Files storage? [y/N] ")
    if answer.strip().lower() != "y":
        print("Cancelled.")
        return
    for f in files:
        try:
            client.files.delete(name=f.name)
            print(f"Deleted {f.display_name or ''} ({f.name})")
        except Exception as e:
            print(f"Could not delete {f.display_name or ''} ({f.name}): {e}")


# =========================================================================
#  SECTION 4 -- MAIN
# =========================================================================

def main() -> int:
    args = parse_args()

    if args.list_types:
        list_types()
        return 0

    if args.purge_all:
        purge_workspace(make_client())
        return 0

    if not args.dir.is_dir():
        sys.exit(f"ERROR: folder not found: {args.dir}")

    # CLI files (if any) replace INPUT_FILES for this run
    cli_entries = list(args.files) + [(name, "container") for name in args.container]
    files = resolve_files(cli_entries or INPUT_FILES, args.dir)
    use_code_execution = code_exec_enabled(args.code_exec, files)
    show_routing(files, use_code_execution)

    if args.dry_run:
        print("Dry run -- nothing converted or uploaded.")
        return 0

    prompt = resolve_prompt(args)
    client = make_client()
    output_path = args.dir / args.output
    mode = "w" if (args.overwrite or not APPEND_MODE) else "a"
    thinking = None if args.thinking == "none" else args.thinking

    uploaded: list[UploadedFile] = []
    generated: list[tuple[bytes, str]] = []
    workdir = tempfile.TemporaryDirectory(prefix="gemini_query_")
    try:
        upload_files(client, files, Path(workdir.name), uploaded)
        contents = build_contents(uploaded, prompt)

        with open(output_path, mode, encoding="utf-8") as out:
            out.write(
                f"\n\n{'=' * 70}\n"
                f" Run: {datetime.now():%Y-%m-%d %H:%M:%S}  |  Model: {args.model}\n"
                f" Files: {', '.join(f'{f.path.name} [{f.route}]' for f in files)}\n"
                f"{'=' * 70}\n\n"
            )
            run_query(
                client, contents, out, generated,
                model=args.model,
                max_tokens=args.max_tokens,
                max_continuations=args.max_continuations,
                system=args.system,
                use_code_execution=use_code_execution,
                thinking_level=thinking,
            )
        print(f"Response written to {output_path}")
        return 0

    except errors.APIError as e:
        print(f"API ERROR: {e}", file=sys.stderr)
        return 1
    except RuntimeError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\nInterrupted -- partial output (if any) was saved.", file=sys.stderr)
        return 130
    finally:
        if generated and DOWNLOAD_GENERATED_FILES:      # saved even after errors
            save_generated_files(generated, args.dir, GENERATED_IMAGE_FORMAT)
        delete_uploads(client, uploaded)                # always clean up
        workdir.cleanup()


if __name__ == "__main__":
    sys.exit(main())

